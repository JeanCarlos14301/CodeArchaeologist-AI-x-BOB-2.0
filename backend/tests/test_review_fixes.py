"""Regression tests for the fixes from the final project review.

- The public showcase reuses one job instead of queuing a pytest run per click.
- The reference comparison only speaks about a reference cut that really ran.
- The live console, the dossier and the memo use the same first-cut PERT (the recommended route).
- The board memo states which route to migrate first.
- A modernization implementation stops at its step and bobcoin limits.
"""

import json
import time
from collections.abc import Iterator
from pathlib import Path

import pytest
from docx import Document
from fastapi.testclient import TestClient

from app.adapters.bob_adapter import BobResult, BobStats
from app.contracts.schema_v1 import MigrationRecommendation, RouteCandidate
from app.main import create_app
from app.modernization.implement import run_steps
from app.modernization.models import Plan
from app.pipeline.activity import EventLog
from app.pipeline.evidence_audit import run_evidence_audit
from app.pipeline.migration_ranking import compare_with_reference

REPO_ROOT = Path(__file__).resolve().parents[2]
SAMPLE = REPO_ROOT / "samples" / "facturaya-v1"
FIXTURE = REPO_ROOT / "contracts" / "fixtures" / "bob-evidence-auditor-facturaya.json"


@pytest.fixture
def client(tmp_path: Path) -> Iterator[TestClient]:
    app = create_app(artifacts_dir=tmp_path / "artifacts", frontend_dist=tmp_path / "no-dist")
    with TestClient(app) as test_client:
        yield test_client


def _wait_done(client: TestClient, job_id: str) -> dict:
    deadline = time.monotonic() + 90
    while time.monotonic() < deadline:
        body = client.get(f"/api/audits/{job_id}").json()
        if body["job"]["status"] in {"done", "failed"}:
            return body
        time.sleep(0.05)
    raise AssertionError("the job did not finish in time")


def _open_showcase(client: TestClient) -> str:
    response = client.post("/api/audits", json={"sample": "facturaya-v1", "execution_mode": "imported"})
    assert response.status_code == 202
    return response.json()["id"]


# --- 1. Showcase -------------------------------------------------------------------------

def test_showcase_is_reused_instead_of_queuing_a_new_run(client: TestClient) -> None:
    first = _open_showcase(client)
    assert _open_showcase(client) == first  # while it is still running
    assert _wait_done(client, first)["job"]["status"] == "done"
    assert _open_showcase(client) == first  # and once it is done
    showcase_jobs = [job for job in client.get("/api/audits").json() if job["execution_mode"] == "imported"]
    assert len(showcase_jobs) == 1


def test_a_failed_showcase_is_replaced(client: TestClient) -> None:
    first = _open_showcase(client)
    _wait_done(client, first)
    client.app.state.audit_service.store.update(first, status="failed", error="simulado")
    assert _open_showcase(client) != first


# --- 2. Reference comparison ---------------------------------------------------------------

def _candidate(endpoint: str, score: float) -> RouteCandidate:
    method, rule = endpoint.split(" ", 1)
    return RouteCandidate(endpoint=endpoint, http_methods=[method], rule=rule, function_name="f", file_path="app.py",
                          line_start=1, line_end=2, value=1.0, risk=1.0, testability=1.0, score=score, formula="-")


def test_reference_comparison_requires_an_executed_reference_cut() -> None:
    recommended = _candidate("GET /invoices/<int:invoice_id>", 2.0)
    recommendation = MigrationRecommendation(recommended=recommended, candidates=[recommended])

    assert compare_with_reference(recommendation, "GET /invoices/{id}", "not_run").reference_comparison is None
    matched = compare_with_reference(recommendation, "GET /invoices/{id}", "passed").reference_comparison
    assert matched and "same endpoint" in matched and "passed" in matched


def test_reference_comparison_reports_a_different_engine_pick() -> None:
    recommended = _candidate("GET /customers", 3.0)
    reference = _candidate("GET /invoices/<int:invoice_id>", 2.0)
    recommendation = MigrationRecommendation(recommended=recommended, candidates=[recommended, reference])

    text = compare_with_reference(recommendation, "GET /invoices/{id}", "passed").reference_comparison

    assert text and "GET /customers" in text and "GET /invoices/<int:invoice_id>" in text


def test_uploads_never_claim_a_laboratory_cut(tmp_path: Path) -> None:
    dossier = run_evidence_audit(SAMPLE, tmp_path / "job", imported_result=FIXTURE, execute_reference_cut=False)
    assert dossier.migration is not None and dossier.migration.status == "not_run"
    assert dossier.recommendation is not None and dossier.recommendation.reference_comparison is None
    # The not-run result names the engine's pick, not a route hardcoded for the sample.
    assert dossier.migration.endpoint == dossier.recommendation.recommended.endpoint


# --- 3. One PERT everywhere ----------------------------------------------------------------

def test_live_console_and_dossier_report_the_same_first_cut_pert(tmp_path: Path) -> None:
    events = EventLog(tmp_path / "job" / "events.jsonl")
    dossier = run_evidence_audit(SAMPLE, tmp_path / "job", imported_result=FIXTURE, events=events)
    lines = (tmp_path / "job" / "events.jsonl").read_text(encoding="utf-8").splitlines()
    summary = next(json.loads(line) for line in lines if '"evidence.summary"' in line)

    assert dossier.recommendation is not None and dossier.first_cut_pert == dossier.recommendation.first_cut_pert
    assert summary["data"]["pert_expected_days"] == dossier.first_cut_pert.expected_days


# --- 4. Memo -------------------------------------------------------------------------------

def test_memo_states_which_route_to_migrate_first(tmp_path: Path) -> None:
    dossier = run_evidence_audit(SAMPLE, tmp_path / "job", imported_result=FIXTURE, execute_reference_cut=True)
    memo = Document(tmp_path / "job" / "board_memo.docx")
    text = "\n".join(paragraph.text for paragraph in memo.paragraphs)
    tables = "\n".join(cell.text for table in memo.tables for row in table.rows for cell in row.cells)
    recommended = dossier.recommendation.recommended

    assert "3 Migration recommendation" in text
    assert f"Authorize the first migration cut on {recommended.endpoint}" in text
    assert recommended.endpoint in tables and "Roadmap by waves" in text
    assert "finding with the highest measured risk" not in text


# --- 5. Modernization limits ---------------------------------------------------------------

def _plan(steps: int) -> Plan:
    return Plan.model_validate({
        "summary": "Test plan with chained steps.",
        "rollback": "The original is kept.",
        "steps": [
            {"id": f"S{i}", "title": f"Step {i}", "kind": "code", "why": "Reason for the test step.",
             "depends_on": [f"S{i - 1}"] if i > 1 else [], "files": [{"path": f"m{i}.py", "action": "create"}],
             "risk": "low", "complexity": "low", "validation": "Review.", "changes": "Create the step's module."}
            for i in range(1, steps + 1)
        ],
    })


class _CostlyBob:
    def __init__(self, cost: float) -> None:
        self.cost = cost
        self.calls = 0

    def run(self, mode: str, prompt: str) -> BobResult:
        self.calls += 1
        return BobResult(mode=mode, status="success", last_message='{"summary": "hecho"}',
                         stats=BobStats(task_id="t", duration_ms=1, session_costs=self.cost), execution_mode="live")


def test_implementation_stops_at_the_step_limit(tmp_path: Path) -> None:
    bob = _CostlyBob(0.1)
    runs, _ = run_steps(bob, _plan(4), "flask -> fastapi", tmp_path, lambda *_: None, max_total_cost=100, max_steps=2)

    assert bob.calls == 2
    assert [run.status for run in runs] == ["done", "done", "skipped", "skipped"]
    assert "at most 2 steps" in runs[2].note


def test_implementation_stops_at_the_bobcoin_budget(tmp_path: Path) -> None:
    bob = _CostlyBob(2.0)
    runs, total = run_steps(bob, _plan(4), "flask -> fastapi", tmp_path, lambda *_: None, max_total_cost=3.0, max_steps=10)

    # Budget checked before each step: 0 -> run (2.0), 2.0 < 3 -> run (4.0), then stop.
    assert bob.calls == 2 and total == 4.0
    assert [run.status for run in runs] == ["done", "done", "skipped", "skipped"]
    assert "3 bobcoin budget" in runs[2].note
