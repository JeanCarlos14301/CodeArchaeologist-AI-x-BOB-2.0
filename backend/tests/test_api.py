"""Pruebas de integración de la API (modos example e imported; no gastan bobcoins)."""

import time
from collections.abc import Iterator
from pathlib import Path

import pytest
from docx import Document
from fastapi.testclient import TestClient

from app.jobs.service import IMPORTED_RECORDED_AT
from app.main import create_app

POLL_TIMEOUT_S = 60
POLL_INTERVAL_S = 0.05


@pytest.fixture
def client(tmp_path: Path) -> Iterator[TestClient]:
    app = create_app(artifacts_dir=tmp_path / "artifacts", frontend_dist=tmp_path / "no-dist")
    with TestClient(app) as test_client:
        yield test_client


def _wait_done(client: TestClient, job_id: str) -> dict:
    deadline = time.monotonic() + POLL_TIMEOUT_S
    while time.monotonic() < deadline:
        body = client.get(f"/api/audits/{job_id}").json()
        if body["job"]["status"] in {"done", "failed"}:
            return body
        time.sleep(POLL_INTERVAL_S)
    raise AssertionError("el job no terminó a tiempo")


def test_health(client: TestClient) -> None:
    assert client.get("/health").json()["status"] == "ok"


def test_bob_status_lists_project_assets(client: TestClient) -> None:
    body = client.get("/api/bob/status").json()
    assert "evidence-auditor" in body["custom_modes"]
    assert "legacy-sql-auditor" in body["subagents"]
    assert "legacy-audit-workflow" in body["skills"]


def test_samples(client: TestClient) -> None:
    assert client.get("/api/samples").json() == [{"id": "facturaya-v1", "name": "facturaya-v1"}]


@pytest.mark.parametrize("mode", ["example", "imported"])
def test_audit_completes_and_serves_dossier(client: TestClient, mode: str) -> None:
    response = client.post("/api/audits", json={"sample": "facturaya-v1", "execution_mode": mode})
    assert response.status_code == 202
    body = _wait_done(client, response.json()["id"])
    assert body["job"]["status"] == "done", body["job"]["error"]
    dossier = body["dossier"]
    assert dossier["execution_mode"] == mode
    assert dossier["stats"]["findings_validated"] >= 6
    assert client.get("/api/audits").json()[0]["id"] == body["job"]["id"]
    if mode == "imported":
        assert dossier["generated_at"] == IMPORTED_RECORDED_AT["facturaya-v1"]
        assert dossier["job_id"] == body["job"]["id"]
        assert len(dossier["source_sha256"]) == 64
        assert dossier["risk_matrix"] and dossier["first_cut_pert"]
        assert dossier["migration"]["status"] == "passed"
        migration = client.get(f"/api/audits/{body['job']['id']}/migration")
        assert migration.status_code == 200
        migration_body = migration.json()
        assert migration_body["legacy_code"] and migration_body["modern_code"]
        assert len(migration_body["result"]["tests"]) == 6
        assert client.get(f"/api/audits/{body['job']['id']}/files/migration.diff").status_code == 200
        memo_response = client.get(f"/api/audits/{body['job']['id']}/files/board_memo.docx")
        assert memo_response.status_code == 200
        memo_path = Path(client.app.state.audit_service.job_dir(body["job"]["id"])) / "board_memo.docx"
        memo = Document(memo_path)
        text = "\n".join(paragraph.text for paragraph in memo.paragraphs)
        tables = "\n".join(cell.text for table in memo.tables for row in table.rows for cell in row.cells)
        assert body["job"]["id"] in tables
        assert dossier["source_sha256"] in tables
        assert "LEGACYLENS" not in (text + tables).upper()
        assert "CBRS" not in (text + tables).upper()
        assert "6 pruebas pasaron y 0 fallaron" in (text + tables)


def test_source_viewer_returns_cited_lines(client: TestClient) -> None:
    job_id = client.post("/api/audits", json={"sample": "facturaya-v1", "execution_mode": "example"}).json()["id"]
    _wait_done(client, job_id)
    excerpt = client.get(f"/api/audits/{job_id}/source", params={"path": "app.py", "start": 78, "end": 79}).json()
    assert excerpt["start"] == 78
    assert "SELECT id, number" in excerpt["lines"][0]["text"]


@pytest.mark.parametrize("path", ["../../../.env", "/etc/passwd", ".bob/custom_modes.yaml", "nope.py"])
def test_source_viewer_rejects_paths_outside_workspace(client: TestClient, path: str) -> None:
    job_id = client.post("/api/audits", json={"sample": "facturaya-v1", "execution_mode": "example"}).json()["id"]
    _wait_done(client, job_id)
    assert client.get(f"/api/audits/{job_id}/source", params={"path": path}).status_code == 404


def test_unknown_sample_is_404(client: TestClient) -> None:
    assert client.post("/api/audits", json={"sample": "otro", "execution_mode": "example"}).status_code == 404


def test_invalid_sample_name_is_422(client: TestClient) -> None:
    assert client.post("/api/audits", json={"sample": "../etc", "execution_mode": "example"}).status_code == 422


def test_second_live_audit_is_rejected_while_one_is_active(client: TestClient, monkeypatch: pytest.MonkeyPatch) -> None:
    # Token explícito: la prueba no depende del LIVE_AUDIT_TOKEN que tenga el .env de quien la ejecuta.
    monkeypatch.setenv("LIVE_AUDIT_TOKEN", "token-de-prueba")
    client.app.state.audit_service.store.create("facturaya-v1", "live")
    response = client.post("/api/audits", json={"sample": "facturaya-v1", "execution_mode": "live"},
                           headers={"X-Live-Token": "token-de-prueba"})
    assert response.status_code == 409


def test_unknown_job_is_404(client: TestClient) -> None:
    assert client.get("/api/audits/nope").status_code == 404


@pytest.mark.parametrize("token", [None, "incorrecto"])
def test_live_requires_token_when_configured(
    client: TestClient, monkeypatch: pytest.MonkeyPatch, token: str | None
) -> None:
    monkeypatch.setenv("LIVE_AUDIT_TOKEN", "secreto-de-prueba")
    headers = {"X-Live-Token": token} if token else {}
    response = client.post("/api/audits", json={"sample": "facturaya-v1", "execution_mode": "live"},
                           headers=headers)
    assert response.status_code == 403
    assert client.get("/api/bob/status").json()["live_requires_token"] is True


def test_live_token_is_accepted_when_correct(client: TestClient, monkeypatch: pytest.MonkeyPatch) -> None:
    monkeypatch.setenv("LIVE_AUDIT_TOKEN", "secreto-de-prueba")
    # Un live ya activo hace que la petición válida responda 409 sin invocar Bob:
    # así se prueba que el token pasó la puerta sin gastar bobcoins.
    client.app.state.audit_service.store.create("facturaya-v1", "live")
    response = client.post("/api/audits", json={"sample": "facturaya-v1", "execution_mode": "live"},
                           headers={"X-Live-Token": "secreto-de-prueba"})
    assert response.status_code == 409


def test_token_never_required_for_example_mode(client: TestClient, monkeypatch: pytest.MonkeyPatch) -> None:
    monkeypatch.setenv("LIVE_AUDIT_TOKEN", "secreto-de-prueba")
    response = client.post("/api/audits", json={"sample": "facturaya-v1", "execution_mode": "example"})
    assert response.status_code == 202


def test_live_is_open_without_configured_token(client: TestClient, monkeypatch: pytest.MonkeyPatch) -> None:
    monkeypatch.delenv("LIVE_AUDIT_TOKEN", raising=False)
    assert client.get("/api/bob/status").json()["live_requires_token"] is False
