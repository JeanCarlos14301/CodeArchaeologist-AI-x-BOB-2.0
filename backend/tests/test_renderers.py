"""Pruebas de los Renderizadores de Artefactos DOCX, HTML y PPTX (D-05, D-08, D-09)."""

import json
from pathlib import Path
from docx import Document
from pptx import Presentation
import pytest

from backend.app.models import DossierResult
from backend.app.renderers.docx_renderer import render_dossier_to_docx
from backend.app.renderers.html_renderer import render_dossier_to_html
from backend.app.renderers.pptx_renderer import create_executive_pptx

BASE_DIR = Path(__file__).resolve().parent.parent.parent


@pytest.fixture
def sample_dossier() -> DossierResult:
    fixture_path = BASE_DIR / "contracts" / "fixtures" / "valid-dossier.json"
    data = json.loads(fixture_path.read_text(encoding="utf-8"))
    return DossierResult.model_validate(data)


def test_docx_renderer_produces_valid_word_document(sample_dossier: DossierResult, tmp_path: Path):
    """Verifica que el renderizador DOCX genere un documento de 7 secciones que abra limpiamente en Word."""
    out_docx = tmp_path / "board_memo_test.docx"
    render_dossier_to_docx(sample_dossier, out_docx)

    assert out_docx.exists()
    assert out_docx.stat().st_size > 5000

    # Abrir con python-docx para verificar que no esté corrupto
    doc = Document(str(out_docx))
    assert len(doc.paragraphs) > 10
    assert len(doc.tables) >= 4  # Hallazgos, Opciones, PERT, etc.

    full_text = "\n".join(p.text for p in doc.paragraphs)
    assert "Resumen Ejecutivo y Decisión Solicitada" in full_text
    assert "Diagnóstico del Sistema y Evidencia Física Verificada" in full_text
    assert "Evaluación de Riesgo y Radio de Explosión" in full_text
    assert "Opciones Arquitectónicas Comparadas" in full_text
    assert "Plan de Fases de Migración con Distribución PERT" in full_text
    assert "Resultados del Primer Corte Probado" in full_text
    assert "Hoja de Ruta para los Próximos 30 Días" in full_text


def test_html_renderer_produces_interactive_report(sample_dossier: DossierResult, tmp_path: Path):
    """Verifica que el renderizador HTML genere un archivo autónomo con diagramas Mermaid."""
    out_html = tmp_path / "report_test.html"
    render_dossier_to_html(sample_dossier, out_html)

    assert out_html.exists()
    content = out_html.read_text(encoding="utf-8")
    assert "<!DOCTYPE html>" in content
    assert "CodeArchaeologist" in content
    assert "erDiagram" in content or "mermaid" in content
    assert "Fidelidad de Evidencia" in content
    assert "EF-1" in content


def test_pptx_renderer_produces_6_slides_presentation(sample_dossier: DossierResult, tmp_path: Path):
    """Verifica que el renderizador PPTX genere una presentación ejecutiva de exactamente 6 diapositivas."""
    out_pptx = tmp_path / "presentation_test.pptx"
    create_executive_pptx(sample_dossier, out_pptx)

    assert out_pptx.exists()
    assert out_pptx.stat().st_size > 10000

    # Abrir con python-pptx
    prs = Presentation(str(out_pptx))
    assert len(prs.slides) == 6, f"Esperadas 6 diapositivas, generadas {len(prs.slides)}"
