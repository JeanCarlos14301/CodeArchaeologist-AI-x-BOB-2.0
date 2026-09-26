"""Memorando DOCX para junta generado exclusivamente desde el expediente medido."""

from pathlib import Path

from docx import Document
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

from app.contracts.schema_v1 import Dossier

BOARD_MEMO_FILE = "board_memo.docx"
NAVY = "17365D"
PALE_BLUE = "EAF2F8"
PALE_GRAY = "F5F6F7"
BORDER = "D9D9D9"


def _shade(cell: object, fill: str) -> None:
    properties = cell._tc.get_or_add_tcPr()  # type: ignore[attr-defined]
    shading = OxmlElement("w:shd")
    shading.set(qn("w:fill"), fill)
    properties.append(shading)


def _borders(table: object) -> None:
    properties = table._tbl.tblPr  # type: ignore[attr-defined]
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        node = OxmlElement(f"w:{edge}")
        node.set(qn("w:val"), "single")
        node.set(qn("w:sz"), "4")
        node.set(qn("w:color"), BORDER)
        borders.append(node)
    properties.append(borders)


def _cell_margin(cell: object, amount: int = 110) -> None:
    properties = cell._tc.get_or_add_tcPr()  # type: ignore[attr-defined]
    margins = OxmlElement("w:tcMar")
    for side in ("top", "left", "bottom", "right"):
        node = OxmlElement(f"w:{side}")
        node.set(qn("w:w"), str(amount))
        node.set(qn("w:type"), "dxa")
        margins.append(node)
    properties.append(margins)


def _format_table(table: object, widths: list[float], margin: int = 110, font_size: float = 8.5) -> None:
    table.alignment = WD_TABLE_ALIGNMENT.CENTER  # type: ignore[attr-defined]
    table.autofit = False  # type: ignore[attr-defined]
    _borders(table)
    for row_index, row in enumerate(table.rows):  # type: ignore[attr-defined]
        for index, cell in enumerate(row.cells):
            cell.width = Inches(widths[index])
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            _cell_margin(cell, margin)
            if row_index == 0:
                _shade(cell, NAVY)
                for run in cell.paragraphs[0].runs:
                    run.font.bold = True
                    run.font.color.rgb = RGBColor(255, 255, 255)
            elif row_index % 2 == 0:
                _shade(cell, PALE_BLUE)
            for paragraph in cell.paragraphs:
                paragraph.paragraph_format.space_after = Pt(0)
                for run in paragraph.runs:
                    run.font.name = "Arial"
                    run.font.size = Pt(font_size)


def _add_heading(document: Document, text: str, level: int = 1) -> object:
    paragraph = document.add_heading(text, level=level)
    paragraph.paragraph_format.keep_with_next = True
    return paragraph


def _configure(document: Document) -> None:
    section = document.sections[0]
    section.top_margin = Inches(0.72)
    section.bottom_margin = Inches(0.72)
    section.left_margin = Inches(0.78)
    section.right_margin = Inches(0.78)
    styles = document.styles
    styles["Normal"].font.name = "Arial"
    styles["Normal"].font.size = Pt(10)
    styles["Normal"].paragraph_format.space_after = Pt(6)
    for name, size in (("Title", 24), ("Heading 1", 15), ("Heading 2", 12)):
        style = styles[name]
        style.font.name = "Arial"
        style.font.size = Pt(size)
        style.font.color.rgb = RGBColor(0, 0, 0)
        if name == "Title" and style.element.pPr is not None:
            border = style.element.pPr.find(qn("w:pBdr"))
            if border is not None:
                style.element.pPr.remove(border)
    styles["Subtitle"].font.name = "Arial"
    styles["Subtitle"].font.color.rgb = RGBColor(0, 0, 0)


def render_board_memo(dossier: Dossier, output_path: Path, job_id: str) -> Path:
    """Crea el memo trazable; no acepta cifras fuera del expediente."""
    document = Document()
    _configure(document)

    title = document.add_paragraph(style="Title")
    title.add_run(f"Memorando de decisión para {dossier.repo_name}")
    title_properties = title._p.get_or_add_pPr()
    border = title_properties.find(qn("w:pBdr"))
    if border is not None:
        title_properties.remove(border)
    subtitle = document.add_paragraph("Auditoría técnica y alcance del primer corte de migración")
    subtitle.style = document.styles["Subtitle"]

    metadata = document.add_table(rows=4, cols=2)
    metadata.cell(0, 0).text, metadata.cell(0, 1).text = "Identificador", job_id
    metadata.cell(1, 0).text, metadata.cell(1, 1).text = "Hash SHA 256", dossier.source_sha256 or "No disponible"
    metadata.cell(2, 0).text, metadata.cell(2, 1).text = "Modo", dossier.execution_mode
    metadata.cell(3, 0).text, metadata.cell(3, 1).text = "Fecha de la auditoría", dossier.generated_at
    _format_table(metadata, [1.55, 5.35])

    validated = dossier.stats.findings_validated
    reported = dossier.stats.findings_reported
    ratio = dossier.stats.evidence_valid_ratio * 100
    opening = document.add_paragraph()
    opening.add_run("Decisión solicitada. ").bold = True
    opening.add_run(
        f"Autorizar la preparación del primer corte sobre el hallazgo de mayor riesgo medido. "
        f"La auditoría validó {validated} de {reported} hallazgos y {ratio:.1f} por ciento de sus referencias de evidencia. "
        "Este memorando usa solo cifras presentes en el expediente y cálculos deterministas; no incorpora narrativa numérica de Bob."
    )

    _add_heading(document, "1 Conclusión ejecutiva")
    severity_counts = {
        severity: sum(1 for finding in dossier.findings if finding.severity == severity)
        for severity in ("critical", "high", "medium", "low")
    }
    document.add_paragraph(
        f"El expediente contiene {severity_counts['critical']} hallazgos críticos, "
        f"{severity_counts['high']} altos, {severity_counts['medium']} medios y "
        f"{severity_counts['low']} bajos. La prioridad se ordena con severidad y llamadores "
        "transitivos medidos en el grafo; el puntaje no es una opinión del modelo."
    )
    if dossier.first_cut_pert:
        pert = dossier.first_cut_pert
        document.add_paragraph(
            f"El primer corte tiene un rango de {pert.optimistic_days:.1f} a "
            f"{pert.pessimistic_days:.1f} días laborables, con valor esperado PERT de "
            f"{pert.expected_days:.2f} días. El rango se debe usar para planificación, no como compromiso de calendario."
        )

    _add_heading(document, "2 Riesgo y radio de impacto")
    document.add_paragraph(
        "La puntuación multiplica el peso de severidad por uno más el número de funciones "
        "que llaman directa o indirectamente a la función citada. Pesos: crítico 4, alto 3, medio 2 y bajo 1."
    )
    risk_by_id = {item.finding_id: item for item in dossier.risk_matrix}
    risk_table = document.add_table(rows=1, cols=5)
    for index, value in enumerate(("ID", "Severidad", "Origen", "Llamadores", "Puntaje")):
        risk_table.cell(0, index).text = value
    ordered_findings = sorted(
        dossier.findings,
        key=lambda finding: risk_by_id.get(finding.id).score if finding.id in risk_by_id else 0,
        reverse=True,
    )
    for finding in ordered_findings:
        metric = risk_by_id.get(finding.id)
        cells = risk_table.add_row().cells
        values = (
            finding.id,
            finding.severity,
            str(metric.origin_functions if metric else 0),
            str(metric.impacted_callers if metric else 0),
            str(metric.score if metric else 0),
        )
        for index, value in enumerate(values):
            cells[index].text = value
            cells[index].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
    _format_table(risk_table, [0.65, 1.15, 1.1, 1.35, 1.0], margin=55, font_size=8)

    effort_heading = _add_heading(document, "3 Alcance y estimación del primer corte")
    effort_heading.paragraph_format.page_break_before = True  # type: ignore[attr-defined]
    if dossier.first_cut_pert:
        pert = dossier.first_cut_pert
        inputs = document.add_table(rows=2, cols=4)
        labels = ("Rutas afectadas", "Funciones afectadas", "Líneas citadas", "Complejidad acumulada")
        values = (pert.affected_routes, pert.affected_functions, pert.affected_lines, pert.affected_complexity)
        for index, label in enumerate(labels):
            inputs.cell(0, index).text = label
            inputs.cell(1, index).text = str(values[index])
            inputs.cell(1, index).paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
        _format_table(inputs, [1.72, 1.72, 1.72, 1.72])
        document.add_paragraph(f"Fórmula aplicada: {pert.formula}")
        pert_table = document.add_table(rows=2, cols=5)
        labels = ("Optimista", "Más probable", "Pesimista", "Esperado", "Varianza")
        values = (
            f"{pert.optimistic_days:.1f} d",
            f"{pert.most_likely_days:.1f} d",
            f"{pert.pessimistic_days:.1f} d",
            f"{pert.expected_days:.2f} d",
            f"{pert.variance:.2f}",
        )
        for index, label in enumerate(labels):
            pert_table.cell(0, index).text = label
            pert_table.cell(1, index).text = values[index]
            pert_table.cell(1, index).paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
        _format_table(pert_table, [1.38] * 5)
        _add_heading(document, "Supuestos", level=2)
        for assumption in pert.assumptions:
            document.add_paragraph(assumption, style="List Bullet")
    else:
        document.add_paragraph("No fue posible estimar un primer corte porque no hubo hallazgos validados.")

    _add_heading(document, "4 Evidencia prioritaria")
    for finding in ordered_findings[:6]:
        metric = risk_by_id.get(finding.id)
        paragraph = document.add_paragraph()
        paragraph.paragraph_format.keep_with_next = True
        paragraph.add_run(f"{finding.id}  {finding.title}").bold = True
        paragraph.add_run(
            f"\nSeveridad {finding.severity}. Riesgo calculado {metric.score if metric else 0}. "
            f"{finding.explanation}"
        )
        for evidence in finding.evidence:
            document.add_paragraph(
                f"{evidence.path}:{evidence.line_start}-{evidence.line_end}",
                style="List Bullet",
            )

    _add_heading(document, "5 Resultado del primer corte")
    if dossier.migration is None:
        document.add_paragraph("El expediente no contiene un resultado de migración.")
    elif dossier.migration.status == "not_run":
        document.add_paragraph(f"No ejecutada. {dossier.migration.reason or ''}")
    else:
        passed = sum(1 for test in dossier.migration.tests if test.status == "passed")
        failed = sum(1 for test in dossier.migration.tests if test.status == "failed")
        document.add_paragraph(
            f"{dossier.migration.implementation_origin} Resultado: {passed} pruebas pasaron y {failed} fallaron."
        )
        test_table = document.add_table(rows=1, cols=3)
        for index, value in enumerate(("Destino", "Prueba", "Resultado")):
            test_table.cell(0, index).text = value
        for test in dossier.migration.tests:
            cells = test_table.add_row().cells
            cells[0].text = test.target
            cells[1].text = test.name
            cells[2].text = test.status
        _format_table(test_table, [1.1, 4.5, 1.3], margin=70, font_size=8)

    _add_heading(document, "6 Trazabilidad y límites")
    document.add_paragraph(
        "Cada hallazgo incluido conserva archivo, líneas y fragmento literal. El hash identifica el contenido analizado. "
        "Los puntajes de riesgo provienen del grafo estático y el esfuerzo proviene de los cuatro insumos mostrados. "
        "La estimación no incluye aprobaciones externas, esperas de despliegue ni trabajo fuera del primer corte."
    )

    output_path.parent.mkdir(parents=True, exist_ok=True)
    document.save(output_path)
    return output_path
